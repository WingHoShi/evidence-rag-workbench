import asyncio
from pathlib import Path
from uuid import uuid4

import httpx
import pymupdf
import pytest
from docx import Document

from evidence_rag.app import create_app
from evidence_rag.indexing import (
    DocumentParser,
    MultilingualE5EmbeddingEngine,
    ParsedSection,
    QdrantChunkIndex,
    chunk_sections,
)
from evidence_rag.settings import Settings


class FakeEmbeddingEngine:
    dimension = 384

    def __init__(self, model_name: str = "fake-e5-v1") -> None:
        self.model_name = model_name
        self.passages: list[str] = []
        self.queries: list[str] = []

    def embed_passages(self, texts: list[str]) -> list[list[float]]:
        self.passages.extend(texts)
        return [self._vector_for(text) for text in texts]

    def embed_queries(self, texts: list[str]) -> list[list[float]]:
        self.queries.extend(texts)
        return [self._vector_for(text) for text in texts]

    def _vector_for(self, text: str) -> list[float]:
        normalized = text.lower()
        first_dimension = 1.0 if "state" in normalized or "状态" in text else 0.0
        second_dimension = 1.0 - first_dimension
        return [first_dimension, second_dimension] + [0.0] * (self.dimension - 2)


class CloseRecordingClient:
    def __init__(self) -> None:
        self.close_calls = 0

    def close(self) -> None:
        self.close_calls += 1


def request(app, method: str, url: str, **kwargs) -> httpx.Response:
    async def send() -> httpx.Response:
        transport = httpx.ASGITransport(app=app)
        async with httpx.AsyncClient(
            transport=transport,
            base_url="http://testserver",
        ) as client:
            return await client.request(method, url, **kwargs)

    return asyncio.run(send())


def create_knowledge_base(app, name: str) -> dict:
    response = request(app, "POST", "/api/knowledge-bases", json={"name": name})
    assert response.status_code == 201, response.text
    return response.json()


def upload(app, knowledge_base_id: str, filename: str, content: bytes) -> dict:
    response = request(
        app,
        "POST",
        f"/api/knowledge-bases/{knowledge_base_id}/documents",
        files={"file": (filename, content, "application/octet-stream")},
    )
    assert response.status_code == 201, response.text
    return response.json()


@pytest.fixture
def indexed_app(tmp_path):
    settings = Settings(
        data_dir=tmp_path / "data",
        qdrant_path=tmp_path / "data" / "qdrant",
        max_file_bytes=1024 * 1024,
        max_knowledge_base_bytes=5 * 1024 * 1024,
    )
    index = QdrantChunkIndex(
        settings,
        collection_name=f"test_chunks_{uuid4().hex}",
    )
    embeddings = FakeEmbeddingEngine()
    return create_app(settings, vector_index=index, embedding_engine=embeddings)


def test_parser_preserves_markdown_txt_docx_and_pdf_locations(tmp_path) -> None:
    markdown = tmp_path / "guide.md"
    markdown.write_text("# State\nCheckpoints persist graph state.\n\n## Memory\nMemory is separate.", encoding="utf-8")
    text = tmp_path / "notes.txt"
    text.write_text("First paragraph.\n\nSecond paragraph.", encoding="utf-8")
    docx_path = tmp_path / "guide.docx"
    docx = Document()
    docx.add_heading("Tools", level=1)
    docx.add_paragraph("Tools call external functions.")
    docx.save(docx_path)
    pdf_path = tmp_path / "guide.pdf"
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "State persistence stores checkpoints.")
    pdf.save(pdf_path)
    pdf.close()

    parser = DocumentParser()
    markdown_sections = parser.parse(markdown)
    text_sections = parser.parse(text)
    docx_sections = parser.parse(docx_path)
    pdf_sections = parser.parse(pdf_path)

    assert markdown_sections[0].heading == "State"
    assert markdown_sections[0].locator == "heading: State"
    assert text_sections[0].locator == "paragraph 1"
    assert docx_sections[0].heading == "Tools"
    assert pdf_sections[0].locator == "page 1"


def test_chunking_preserves_section_heading_and_locator() -> None:
    sections = [
        ParsedSection(
            text="state " * 80,
            heading="State",
            locator="heading: State",
        )
    ]

    chunks = chunk_sections(sections, chunk_size=80, overlap=16)

    assert len(chunks) > 1
    assert {chunk.heading for chunk in chunks} == {"State"}
    assert {chunk.locator for chunk in chunks} == {"heading: State"}
    assert all("state" in chunk.text for chunk in chunks)


def test_e5_uses_required_query_and_passage_prefixes() -> None:
    assert MultilingualE5EmbeddingEngine.prefix_passage("hello") == "passage: hello"
    assert MultilingualE5EmbeddingEngine.prefix_query("你好") == "query: 你好"


def test_qdrant_chunk_index_closes_the_created_client() -> None:
    index = QdrantChunkIndex(Settings())
    client = CloseRecordingClient()
    index._client = client

    index.close()

    assert client.close_calls == 1
    assert index._client is None


def test_indexes_markdown_and_retrieves_english_evidence_for_chinese_query(indexed_app) -> None:
    knowledge_base = create_knowledge_base(indexed_app, "LangGraph")
    document = upload(
        indexed_app,
        knowledge_base["id"],
        "state.md",
        b"# State\nState persistence stores graph checkpoints for later recovery.",
    )

    indexed = request(
        indexed_app,
        "POST",
        f"/api/knowledge-bases/{knowledge_base['id']}/documents/{document['id']}/index",
    )
    results = indexed_app.state.indexing_service.search(knowledge_base["id"], "状态如何保存？")

    assert indexed.status_code == 200, indexed.text
    assert indexed.json()["status"] == "completed"
    assert results[0]["original_filename"] == "state.md"
    assert results[0]["locator"] == "heading: State"
    assert "checkpoints" in results[0]["text"]
    assert indexed_app.state.indexing_service.embedding_engine.passages
    assert indexed_app.state.indexing_service.embedding_engine.queries == ["状态如何保存？"]


def test_qdrant_filter_prevents_cross_knowledge_base_retrieval(indexed_app) -> None:
    first = create_knowledge_base(indexed_app, "A")
    second = create_knowledge_base(indexed_app, "B")
    first_document = upload(indexed_app, first["id"], "a.md", b"# State\nState is from A.")
    second_document = upload(indexed_app, second["id"], "b.md", b"# State\nState is from B.")

    assert request(indexed_app, "POST", f"/api/knowledge-bases/{first['id']}/documents/{first_document['id']}/index").status_code == 200
    assert request(indexed_app, "POST", f"/api/knowledge-bases/{second['id']}/documents/{second_document['id']}/index").status_code == 200

    results = indexed_app.state.indexing_service.search(first["id"], "状态")

    assert [result["original_filename"] for result in results] == ["a.md"]


def test_pdf_page_is_retained_as_citation_locator(indexed_app, tmp_path) -> None:
    source = tmp_path / "source.pdf"
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "State persistence stores checkpoints.")
    pdf.save(source)
    pdf.close()
    knowledge_base = create_knowledge_base(indexed_app, "PDF")
    document = upload(indexed_app, knowledge_base["id"], "source.pdf", source.read_bytes())

    response = request(
        indexed_app,
        "POST",
        f"/api/knowledge-bases/{knowledge_base['id']}/documents/{document['id']}/index",
    )
    results = indexed_app.state.indexing_service.search(knowledge_base["id"], "状态")

    assert response.status_code == 200, response.text
    assert results[0]["locator"] == "page 1"
    assert "checkpoints" in results[0]["text"]


def test_empty_pdf_fails_task_instead_of_claiming_it_was_indexed(indexed_app, tmp_path) -> None:
    source = tmp_path / "empty.pdf"
    pdf = pymupdf.open()
    pdf.new_page()
    pdf.save(source)
    pdf.close()
    knowledge_base = create_knowledge_base(indexed_app, "扫描件")
    document = upload(indexed_app, knowledge_base["id"], "empty.pdf", source.read_bytes())

    response = request(
        indexed_app,
        "POST",
        f"/api/knowledge-bases/{knowledge_base['id']}/documents/{document['id']}/index",
    )
    tasks = request(
        indexed_app,
        "GET",
        f"/api/knowledge-bases/{knowledge_base['id']}/documents/{document['id']}/tasks",
    )

    assert response.status_code == 422
    assert tasks.json()[-1]["status"] == "failed"


def test_embedding_model_change_requires_and_then_runs_rebuild(indexed_app) -> None:
    knowledge_base = create_knowledge_base(indexed_app, "重建")
    first = upload(indexed_app, knowledge_base["id"], "first.md", b"# State\nFirst state document.")
    assert request(indexed_app, "POST", f"/api/knowledge-bases/{knowledge_base['id']}/documents/{first['id']}/index").status_code == 200
    indexed_app.state.indexing_service.embedding_engine = FakeEmbeddingEngine("fake-e5-v2")
    second = upload(indexed_app, knowledge_base["id"], "second.md", b"# State\nSecond state document.")

    changed_model = request(
        indexed_app,
        "POST",
        f"/api/knowledge-bases/{knowledge_base['id']}/documents/{second['id']}/index",
    )
    rebuilt = request(indexed_app, "POST", f"/api/knowledge-bases/{knowledge_base['id']}/rebuild-index")
    status = request(indexed_app, "GET", f"/api/knowledge-bases/{knowledge_base['id']}/index-status")

    assert changed_model.status_code == 409
    assert "重建" in changed_model.json()["detail"]
    assert rebuilt.status_code == 200, rebuilt.text
    assert status.json()["embedding_model"] == "fake-e5-v2"
    assert status.json()["status"] == "ready"
